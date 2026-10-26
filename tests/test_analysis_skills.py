"""Estimações reais sobre dados sintéticos; não acessam IA nem banco externo."""
import copy
import io

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook
from test_agent import FakeMCP, new_run

from core.models import AgentError
from tools.data_tools import Toolbox, definitions
from tools.forecasting import ForecastRequest, estimate_forecast


def history(model='VAR', size=360):
    rng = np.random.default_rng(31 if model == 'VECM' else 42)
    if model == 'VECM':
        level = np.cumsum(rng.normal(size=size))
        values = np.column_stack([level, level + rng.normal(scale=.4, size=size)])
    elif model in {'ARCH', 'GARCH'}:
        values = np.zeros((size, 1))
        variance = 1.
        for index in range(1, size):
            variance = .3 + .2 * values[index - 1, 0] ** 2 + .5 * variance
            values[index, 0] = rng.normal() * np.sqrt(variance)
    else:
        values = np.zeros((size, 2))
        for index in range(1, size):
            values[index] = np.array([[.45, .1], [.05, .35]]) @ values[index - 1] + rng.normal(size=2)
    dates = pd.date_range('1990-01-01', periods=size, freq='MS')
    return [{'data': date.isoformat(), **{f'serie{j}': float(value) for j, value in enumerate(row)}}
            for date, row in zip(dates, values)]


@pytest.mark.parametrize('model', ['VAR', 'VECM', 'ARCH', 'GARCH'])
async def test_forecast_models_return_finite_future_with_backtest_and_lineage(settings, store, model):
    rows = history(model)
    original = store.save_dataset(rows, {'kind': 'postgresql', 'database': 'synap'})
    tools = Toolbox(settings, store, FakeMCP(), new_run())
    result = await tools.call('forecast_series', {
        'dataset_id': original['id'], 'model': model, 'date_column': 'data',
        'value_columns': ['serie0'] if model in {'ARCH', 'GARCH'} else ['serie0', 'serie1'],
        'horizon': 3, 'backtest_steps': 6,
    })
    projected = store.get_dataset(result['dataset_id'])
    assert result['model'] == model
    assert result['backtest']['steps'] == 6
    assert projected['source']['parent_dataset_id'] == original['id']
    assert projected['source']['original_source']['database'] == 'synap'
    assert projected['source']['forecast_only'] is True
    assert projected['source']['analysis_warnings']
    assert store.get_dataset(original['id'])['rows'] == rows
    assert len(projected['rows']) == (3 if model in {'ARCH', 'GARCH'} else 6)
    for row in projected['rows']:
        assert pd.to_datetime(row['data'], utc=True) > pd.to_datetime(rows[-1]['data'], utc=True)
        assert np.isfinite(row['projecao'])
        assert row['limite_inferior_95'] <= row['projecao'] <= row['limite_superior_95']
        if model in {'ARCH', 'GARCH'}:
            assert row['volatilidade'] > 0
            assert row['volatilidade'] ** 2 == pytest.approx(row['variancia'])
    if model == 'VECM':
        assert result['diagnostics']['cointegration_rank_training'] == 1
    if model == 'VAR':
        transformed = await tools.call('transform_dataset', {
            'dataset_id': projected['id'], 'operations': [{'op': 'select', 'columns': ['data', 'projecao']}],
        })
        assert any('Backtest' in warning for warning in transformed['warnings'])
        report = await tools.call('create_report', {'title': 'Projeção', 'body': 'Estimativas futuras.',
                                                   'dataset_ids': [transformed['dataset_id']]})
        for artifact in report['artifacts']:
            content = (store.artifacts_dir / artifact['filename']).read_text(encoding='utf-8')
            assert 'Backtest' in content and 'não são garantias' in content


@pytest.mark.parametrize('problem', ['short', 'duplicate_date', 'gap', 'null', 'numeric_text', 'constant', 'missing_column'])
def test_forecasts_refuse_invalid_history_without_imputation(problem):
    rows = history()
    if problem == 'short':
        rows = rows[:20]
    elif problem == 'duplicate_date':
        rows[1]['data'] = rows[0]['data']
    elif problem == 'gap':
        rows.pop(30)
    elif problem == 'null':
        rows[10]['serie0'] = None
    elif problem == 'numeric_text':
        for row in rows:
            row['serie0'] = str(row['serie0'])
    elif problem == 'constant':
        for row in rows:
            row['serie0'] = 1.
    elif problem == 'missing_column':
        rows = [{'data': row['data'], 'serie0': row['serie0']} for row in rows]
    request = ForecastRequest(dataset_id='synthetic', date_column='data', value_columns=['serie0', 'serie1'], model='VAR')
    with pytest.raises(AgentError):
        estimate_forecast(request, {'rows': rows})


def test_forecasting_does_not_silently_replace_incompatible_model():
    request = ForecastRequest(dataset_id='synthetic', date_column='data', value_columns=['serie0', 'serie1'], model='VECM')
    with pytest.raises(AgentError, match=r'I\(1\)'):
        estimate_forecast(request, {'rows': history('VAR')})


def test_holdout_is_not_used_in_stationarity_diagnostics():
    request = ForecastRequest(dataset_id='synthetic', date_column='data', value_columns=['serie0', 'serie1'], model='VAR')
    rows = history()
    initial = estimate_forecast(request, {'rows': rows})
    changed = copy.deepcopy(rows)
    for row in changed[-6:]:
        row['serie0'] += 2
    revised = estimate_forecast(request, {'rows': changed})
    assert initial['diagnostics']['adf_training_pvalues'] == revised['diagnostics']['adf_training_pvalues']
    assert initial['backtest']['mae_by_series'] != revised['backtest']['mae_by_series']


async def test_pattern_finder_returns_evidence_without_mutating_data(settings, store):
    rows = [{'volume': index, 'receita': index * 3, 'atipico': index if index < 49 else 10000} for index in range(50)]
    dataset = store.save_dataset(rows, {'kind': 'postgresql', 'possibly_truncated': True})
    tools = Toolbox(settings, store, FakeMCP(), new_run())
    result = await tools.call('find_patterns', {'dataset_id': dataset['id'], 'columns': ['volume', 'receita', 'atipico']})
    pair = next(pair for pair in result['correlations'] if pair['left'] == 'volume' and pair['right'] == 'receita')
    assert pair['spearman'] == pytest.approx(1)
    assert pair['observations'] == 50
    assert any(item['column'] == 'atipico' and item['count'] == 1 for item in result['outliers'])
    assert result['possibly_truncated'] is True
    assert any('causalidade' in warning for warning in result['warnings'])
    assert store.get_dataset(dataset['id'])['rows'] == rows


async def test_professional_excel_contains_business_sheets_and_chart_without_technical_sources(settings, store):
    rows = [{'mes': 'Janeiro', 'saldo': 1234.5, 'nome': '=HYPERLINK("bad")'},
            {'mes': 'Fevereiro', 'saldo': -150.0, 'nome': 'Normal'}]
    dataset = store.save_dataset(rows, {'kind': 'postgresql', 'database': 'cora', 'table': 'tabela_tecnica'})
    tools = Toolbox(settings, store, FakeMCP(), new_run())
    result = await tools.call('create_excel_report', {
        'title': 'Carteira de crédito', 'summary': 'Comparação mensal dos saldos.', 'period': '2026',
        'dataset_ids': [dataset['id']], 'sheets': [
            {'dataset_id': dataset['id'], 'title': 'Saldos mensais', 'context': 'Saldos de referência por mês.',
             'columns': [{'key': 'mes', 'label': 'Mês'}, {'key': 'saldo', 'label': 'Saldo', 'unit': 'R$'},
                         {'key': 'nome', 'label': 'Observação'}],
             'chart': 'bar', 'category_key': 'mes', 'value_key': 'saldo'}],
    })
    raw = (store.artifacts_dir / result['artifacts'][0]['filename']).read_bytes()
    workbook = load_workbook(io.BytesIO(raw))
    assert workbook.sheetnames == ['Resumo', 'Saldos mensais']
    sheet = workbook['Saldos mensais']
    assert sheet['B6'].value == 1234.5
    assert sheet['B7'].value == -150.
    assert sheet['C6'].value.startswith("'=")
    assert sheet['C6'].data_type == 's'
    assert sheet.freeze_panes == 'A6'
    assert len(sheet._charts) == 1
    text = str([[cell.value for row in sheet for cell in row] for sheet in workbook])
    assert 'Banco de dados interno' in text
    for hidden in ['tabela_tecnica', dataset['id'], "'cora'"]:
        assert hidden not in text


def test_new_tools_are_declarative_and_registered():
    schemas = {item['name']: item['parameters'] for item in definitions()}
    assert {'forecast_series', 'find_patterns', 'create_excel_report'} <= schemas.keys()
    assert schemas['forecast_series']['additionalProperties'] is False


@pytest.mark.parametrize('problem', ['text', 'small', 'missing'])
async def test_pattern_finder_refuses_invalid_measurements(settings, store, problem):
    rows = [{'valor': index} for index in range(30)]
    if problem == 'small':
        rows = rows[:5]
    elif problem == 'text':
        rows = [{'valor': str(row['valor'])} for row in rows]
    dataset = store.save_dataset(rows, {'kind': 'test'})
    tools = Toolbox(settings, store, FakeMCP(), new_run())
    with pytest.raises(AgentError):
        await tools.call('find_patterns', {'dataset_id': dataset['id'], 'columns': ['ausente' if problem == 'missing' else 'valor']})


@pytest.mark.parametrize('problem', ['dataset', 'chart_column', 'chart_text'])
async def test_excel_validation_does_not_publish_invalid_report(settings, store, problem):
    dataset = store.save_dataset([{'mes': 'Jan', 'saldo': '10.50' if problem == 'chart_text' else 10.5}], {'kind': 'test'})
    tools = Toolbox(settings, store, FakeMCP(), new_run())
    with pytest.raises(AgentError):
        await tools.call('create_excel_report', {
            'title': 'Carteira', 'period': '2026', 'summary': 'Saldos mensais.', 'dataset_ids': [dataset['id']],
            'sheets': [{'dataset_id': 'inexistente' if problem == 'dataset' else dataset['id'],
                        'title': 'Saldos', 'context': 'Saldos da carteira no período.',
                        'columns': [{'key': 'mes', 'label': 'Mês'}, {'key': 'saldo', 'label': 'Saldo'}],
                        'chart': 'bar', 'category_key': 'mes', 'value_key': 'ausente' if problem == 'chart_column' else 'saldo'}],
        })
    assert list(store.artifacts_dir.iterdir()) == []
