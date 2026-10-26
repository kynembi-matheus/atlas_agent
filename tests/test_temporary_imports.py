"""Importação com IA simulada, Parquet real e comprovação de isolamento/remoção."""
import asyncio
import io
import json
import threading
import time
from concurrent.futures import ThreadPoolExecutor
from types import SimpleNamespace
from uuid import uuid4
from zipfile import ZIP_DEFLATED, ZIP_STORED, ZipFile

import pyarrow as pa
import pyarrow.parquet as pq
import pytest
import xlwt
from fastapi.testclient import TestClient
from openpyxl import Workbook
from pydantic import SecretStr
from test_agent import ScriptLLM

from api.main import create_app
from core.import_sessions import ImportSessions
from core.models import AgentError
from integrations.file_imports import ImportFailure, office_archive
from integrations.temporary_imports import ImportSchema, infer_schema, normalize_tables, read_tables
from repositories.temporary_store import remove_session_directory
from tools.data_tools import Toolbox


def schema(tables, header_index=0):
    result = []
    for table in tables:
        width = max(map(len, table.rows))
        labels = table.rows[header_index] if header_index is not None else [f'Campo {i}' for i in range(width)]
        columns = []
        for index, name in enumerate(labels):
            dtype = {'Codigo': 'integer', 'Valor': 'decimal', 'Data': 'date', 'Quantidade': 'integer'}.get(name, 'string')
            columns.append({'index': index, 'name': str(name), 'dtype': dtype,
                            'number_format': 'decimal_comma', 'date_format': '%d/%m/%Y'})
        result.append({'sheet_id': table.sheet_id, 'header_index': header_index, 'confidence': .95, 'columns': columns})
    return {'tables': result}


class InferenceLLM:
    def __init__(self, *, malformed=False):
        self.responses = self
        self.calls = []
        self.malformed = malformed

    async def create(self, **kwargs):
        self.calls.append(kwargs)
        samples = json.loads(kwargs['input'][0]['content'])
        tables = [SimpleNamespace(sheet_id=sample['sheet_id'], rows=sample['sample']) for sample in samples]
        arguments = '{' if self.malformed else json.dumps(schema(tables))
        return SimpleNamespace(output=[SimpleNamespace(type='function_call', name='infer_import_schema', arguments=arguments)])


def start(client, previous=None):
    response = client.post('/api/import-sessions', headers={'X-Data-Session': previous} if previous else {})
    assert response.status_code == 201
    return response.json()['session_id']


def upload(client, session_id, name='dados.csv', raw=b'Codigo;Valor;Data;Quantidade\n00123;1.234,56;26/10/2026;2\n00456;10,50;27/10/2026;3'):
    return client.post('/api/imports', headers={'X-Data-Session': session_id}, files={'file': (name, raw)})


def test_csv_inference_preserves_identifiers_and_writes_typed_parquet_only(settings):
    llm = InferenceLLM()
    with TestClient(create_app(settings, mcp=object(), client=llm)) as client:
        session_id = start(client)
        response = upload(client, session_id)
        assert response.status_code == 200, response.text
        data = response.json()['datasets'][0]
        assert data['preview'][0]['Codigo'] == '00123'
        assert data['preview'][0]['Valor'] == 1234.56
        assert data['preview'][0]['Quantidade'] == 2
        assert data['source']['inferred_types']['Codigo'] == 'string'
        session = client.app.state.imports.get(session_id)
        files = list(session.store.root.glob('*.parquet'))
        assert len(files) == 1
        assert not list(session.store.root.glob('*.csv'))
        fields = pq.read_schema(files[0])
        assert pa.types.is_string(fields.field('Codigo').type) or pa.types.is_large_string(fields.field('Codigo').type)
        assert pa.types.is_integer(fields.field('Quantidade').type)
        assert pa.types.is_timestamp(fields.field('Data').type)
        with client.app.state.store.connect() as db:
            assert db.execute('SELECT COUNT(*) FROM datasets').fetchone()[0] == 0
        assert client.get(f'/api/datasets/{data["dataset_id"]}', headers={'X-Data-Session': session_id}).status_code == 200
        assert client.get(f'/api/datasets/{data["dataset_id"]}').status_code == 404
        root = session.store.root
    assert not root.exists()


def test_new_session_deletes_old_parquet_and_isolates_another_session(settings):
    with TestClient(create_app(settings, mcp=object(), client=InferenceLLM())) as client:
        first, second = start(client), start(client)
        first_data = upload(client, first).json()['datasets'][0]
        second_data = upload(client, second).json()['datasets'][0]
        first_root = client.app.state.imports.get(first).store.root
        second_root = client.app.state.imports.get(second).store.root
        assert client.get(f'/api/datasets/{first_data["dataset_id"]}', headers={'X-Data-Session': second}).status_code == 404
        new = start(client, first)
        assert new != first and not first_root.exists() and second_root.exists()
        assert client.get(f'/api/datasets/{first_data["dataset_id"]}', headers={'X-Data-Session': first}).status_code == 410
        assert client.get(f'/api/datasets/{first_data["dataset_id"]}', headers={'X-Data-Session': new}).status_code == 404
        assert client.get(f'/api/datasets/{second_data["dataset_id"]}', headers={'X-Data-Session': second}).status_code == 200


def test_temporary_runs_events_derived_datasets_and_reports_never_reach_sqlite(settings):
    raw = b'Nome;Valor\nCONTEUDO_TEMPORARIO_PRIVADO;25,50'
    tables, _ = read_tables('dados.csv', raw, 100)
    llm = ScriptLLM([('infer_import_schema', schema(tables))])
    with TestClient(create_app(settings, mcp=object(), client=llm)) as client:
        session_id = start(client)
        data = upload(client, session_id, raw=raw).json()['datasets'][0]
        llm.actions = [('inspect_dataset', {'dataset_id': data['dataset_id']}),
                       ('transform_dataset', {'dataset_id': data['dataset_id'], 'operations': [{'op': 'select', 'columns': ['Nome', 'Valor']}]}),
                       ('create_report', {'title': 'Análise', 'body': 'CONTEUDO_TEMPORARIO_PRIVADO', 'dataset_ids': [data['dataset_id']]}),
                       'CONTEUDO_TEMPORARIO_PRIVADO']
        headers = {'X-Data-Session': session_id}
        response = client.post('/api/runs', headers=headers, json={'prompt': 'Analise os dados.', 'dataset_ids': [data['dataset_id']]})
        assert response.status_code == 202, response.text
        run_id = response.json()['id']
        for _ in range(100):
            run = client.get(f'/api/runs/{run_id}', headers=headers).json()
            if run['status'] not in {'running', 'queued'}:
                break
            time.sleep(.01)
        assert run['status'] == 'completed', run
        assert run['answer'] == 'CONTEUDO_TEMPORARIO_PRIVADO'
        assert len(run['artifacts']) == 2
        assert client.get(run['artifacts'][0]['url'], headers=headers).status_code == 200
        assert client.get(run['artifacts'][0]['url']).status_code == 404
        with client.app.state.store.connect() as db:
            for table in ['datasets', 'runs', 'artifacts']:
                assert db.execute(f'SELECT COUNT(*) FROM {table}').fetchone()[0] == 0
        assert client.get(f'/api/runs/{run_id}').status_code == 404
        session = client.app.state.imports.get(session_id)
        assert len(session.store.datasets) == 2
        root = session.store.root
        start(client, session_id)
        assert not root.exists()


def test_xlsx_tabs_are_separate_with_source_rows_and_cached_formulas(settings):
    book = Workbook()
    book.active.title = 'Carteira'
    for sheet in [book.active, book.create_sheet('Cobrança')]:
        sheet.append(['Codigo', 'Valor'])
        sheet.append([123, 50.25])
        sheet['A2'].number_format = '00000'
    buffer = io.BytesIO()
    book.save(buffer)
    with TestClient(create_app(settings, mcp=object(), client=InferenceLLM())) as client:
        session = start(client)
        response = upload(client, session, 'dados.xlsx', buffer.getvalue())
        assert response.status_code == 200, response.text
        datasets = response.json()['datasets']
        assert len(datasets) == 2
        assert {data['source']['sheet'] for data in datasets} == {'Carteira', 'Cobrança'}
        assert all(data['preview'][0]['Codigo'] == '00123' for data in datasets)
        assert all(data['preview'][0]['__linha_origem'] == 2 for data in datasets)
        assert len(list(client.app.state.imports.get(session).store.root.glob('*.parquet'))) == 2


def test_real_xls_values_are_imported_without_running_macros(settings):
    book = xlwt.Workbook()
    sheet = book.add_sheet('Legado')
    sheet.write(0, 0, 'Codigo')
    sheet.write(0, 1, 'Valor')
    sheet.write(1, 0, 123, xlwt.easyxf(num_format_str='00000'))
    sheet.write(1, 1, 25.5)
    buffer = io.BytesIO()
    book.save(buffer)
    with TestClient(create_app(settings, mcp=object(), client=InferenceLLM())) as client:
        session = start(client)
        response = upload(client, session, 'legado.xls', buffer.getvalue())
        assert response.status_code == 200, response.text
        data = response.json()['datasets'][0]
        assert data['preview'][0]['Codigo'] == '00123'
        assert data['preview'][0]['Valor'] == 25.5
        assert data['source']['sheet'] == 'Legado'


def test_inference_failure_saves_nothing_and_unsupported_formats_do_not_call_ai(settings):
    llm = InferenceLLM(malformed=True)
    with TestClient(create_app(settings, mcp=object(), client=llm)) as client:
        session = start(client)
        assert upload(client, session).status_code == 503
        assert list(client.app.state.imports.get(session).store.root.glob('*.parquet')) == []
        assert upload(client, session, 'arquivo.exe', b'bytes').status_code == 415
        assert len(llm.calls) == 1
        assert client.post('/api/imports', files={'file': ('dados.csv', b'A\n1')}).status_code == 409


def test_inference_applies_to_entire_column_without_dropping_incompatible_rows():
    raw = b'Titulo do relatorio\nCodigo;Valor\n00123;1.234,56\n00456;valor indisponivel'
    tables, notes = read_tables('dados.csv', raw, 100)
    specs = ImportSchema.model_validate(schema(tables, header_index=1))
    rows, source = normalize_tables('dados.csv', raw, tables, specs, notes, 100)[0]
    assert len(rows) == 2
    assert rows[0]['Valor'] == '1.234,56'
    assert rows[1]['Valor'] == 'valor indisponivel'
    assert source['header_row'] == 2
    assert source['inferred_types']['Valor'] == 'string'
    assert source['extraction_warnings']


async def test_ai_receives_only_bounded_samples_and_cannot_choose_other_tools():
    raw = ('Codigo;Valor\n' + '\n'.join(f'{i:05};{i},50' for i in range(100))).encode()
    tables, _ = read_tables('dados.csv', raw, 500)
    llm = InferenceLLM()
    await infer_schema(tables, llm, 'fake-model')
    call = llm.calls[0]
    assert call['store'] is False
    assert call['tool_choice'] == {'type': 'function', 'name': 'infer_import_schema'}
    payload = call['input'][0]['content']
    assert len(payload) <= 20000
    assert len(json.loads(payload)[0]['sample']) == 20
    assert '00099' not in payload


async def test_startup_ttl_and_closed_store_cleanup_are_scoped(settings, store):
    root = settings.data_dir.parent / 'temp' / 'import-sessions'
    orphan = root / str(uuid4())
    orphan.mkdir(parents=True)
    (orphan / 'old.parquet').write_bytes(b'synthetic')
    keep = root / 'other-file.txt'
    keep.write_text('preservar', encoding='utf-8')
    registry = ImportSessions(settings, store, object(), object())
    assert not orphan.exists() and keep.exists()
    old, active = registry.create(), registry.create()
    old.touched -= registry.ttl + 1
    await registry.cleanup()
    assert old.id not in registry.sessions and active.id in registry.sessions
    with pytest.raises(AgentError):
        old.store.save_dataset([{'valor': 1}], {})
    await registry.close()
    assert not active.store.root.exists() and keep.exists()


def test_recursive_cleanup_refuses_paths_outside_temp(tmp_path):
    root = tmp_path / 'managed'
    root.mkdir()
    outside = tmp_path / str(uuid4())
    outside.mkdir()
    with pytest.raises(AgentError):
        remove_session_directory(outside, root)
    assert outside.exists()


def test_normalization_rejects_inconsistent_ai_columns():
    raw = b'A;B\n1;2'
    tables, notes = read_tables('dados.csv', raw, 100)
    spec = schema(tables)
    spec['tables'][0]['columns'].pop()
    with pytest.raises(ImportFailure):
        normalize_tables('dados.csv', raw, tables, ImportSchema.model_validate(spec), notes, 100)


def test_reset_while_ai_is_running_does_not_recreate_expired_files(settings):
    started, release = threading.Event(), threading.Event()

    class SlowInference(InferenceLLM):
        async def create(self, **kwargs):
            started.set()
            await asyncio.to_thread(release.wait, 5)
            return await super().create(**kwargs)

    with TestClient(create_app(settings, mcp=object(), client=SlowInference())) as client:
        session = start(client)
        root = client.app.state.imports.get(session).store.root
        with ThreadPoolExecutor(max_workers=1) as pool:
            pending = pool.submit(upload, client, session)
            assert started.wait(3)
            try:
                new = start(client, session)
            finally:
                release.set()
            response = pending.result(timeout=5)
        assert response.status_code == 410
        assert not root.exists()
        assert not list(client.app.state.imports.get(new).store.root.glob('*.parquet'))


def test_import_limits_and_missing_ai_do_not_write_temporary_files(settings):
    settings.max_temporary_upload_bytes = 1024
    llm = InferenceLLM()
    with TestClient(create_app(settings, mcp=object(), client=llm)) as client:
        session = start(client)
        assert upload(client, session, raw=b'A\n' + b'1' * 1025).status_code == 413
        assert llm.calls == []
        assert list(client.app.state.imports.get(session).store.root.glob('*.parquet')) == []
        settings.openai_api_key = SecretStr('')
        assert upload(client, session).status_code == 409
        assert llm.calls == []


async def test_import_list_is_scoped_to_current_session(settings, store):
    registry = ImportSessions(settings, store, object(), object())
    first, second = registry.create(), registry.create()
    data = first.store.save_dataset([{'valor': 1}], {'kind': 'temporary_import', 'filename': 'dados.csv'})
    result = await Toolbox(settings, first.store, object(), None).call('list_imported_datasets', {})
    assert [item['dataset_id'] for item in result['datasets']] == [data['id']]
    assert not (await Toolbox(settings, second.store, object(), None).call('list_imported_datasets', {}))['datasets']
    assert not (await Toolbox(settings, store, object(), None).call('list_imported_datasets', {}))['datasets']
    await registry.close()


async def test_parquet_roundtrip_does_not_round_numeric_values(settings, store):
    registry = ImportSessions(settings, store, object(), object())
    session = registry.create()
    rows = [{'medida': .12345678901234567, 'inteiro': 9223372036854775807}]
    dataset = session.store.save_dataset(rows, {})
    assert session.store.get_dataset(dataset['id'])['rows'] == rows
    await registry.close()


def test_partial_parquet_failure_rolls_back_the_entire_workbook(settings, monkeypatch):
    book = Workbook()
    for sheet in [book.active, book.create_sheet('Segunda')]:
        sheet.append(['Valor'])
        sheet.append([25.5])
    raw = io.BytesIO()
    book.save(raw)
    with TestClient(create_app(settings, mcp=object(), client=InferenceLLM())) as client:
        session = start(client)
        temporary = client.app.state.imports.get(session).store
        save = temporary.save_dataset
        calls = 0

        def fail_second(rows, source):
            nonlocal calls
            calls += 1
            if calls == 2:
                raise AgentError('Falha simulada')
            return save(rows, source)

        monkeypatch.setattr(temporary, 'save_dataset', fail_second)
        response = upload(client, session, 'dados.xlsx', raw.getvalue())
        assert response.status_code == 422
        assert temporary.datasets == {}
        assert list(temporary.root.glob('*.parquet')) == []


def test_session_imports_can_be_reused_without_reattaching_and_stay_temporary(settings):
    raw = b'Nome;Valor\nTemporario;25,50'
    tables, _ = read_tables('dados.csv', raw, 100)
    llm = ScriptLLM([('infer_import_schema', schema(tables))])
    with TestClient(create_app(settings, mcp=object(), client=llm)) as client:
        session_id = start(client)
        data = upload(client, session_id, raw=raw).json()['datasets'][0]
        llm.actions = [('list_imported_datasets', {}), 'Os dados da sessão foram encontrados.']
        headers = {'X-Data-Session': session_id}
        response = client.post('/api/runs', headers=headers, json={'prompt': 'Use meus dados importados.'})
        assert response.status_code == 202
        for _ in range(100):
            run = client.get('/api/runs/' + response.json()['id'], headers=headers).json()
            if run['status'] == 'completed':
                break
            time.sleep(.01)
        assert run['status'] == 'completed'
        found = next(event['result'] for event in run['events'] if event.get('tool') == 'list_imported_datasets' and event['kind'] == 'tool_result')
        assert found['datasets'][0]['dataset_id'] == data['dataset_id']
        assert client.app.state.store.list_runs() == []


def test_temporary_upload_default_is_500_mb_and_accepts_files_above_legacy_limit(settings):
    book = Workbook()
    book.active.append(['Valor'])
    book.active.append([25.5])
    buffer = io.BytesIO()
    book.save(buffer)
    with ZipFile(buffer, 'a') as archive:
        archive.writestr('padding.bin', b'\0' * (6 * 1024 * 1024), compress_type=ZIP_STORED)
    raw = buffer.getvalue()
    assert len(raw) > settings.max_upload_bytes
    with TestClient(create_app(settings, mcp=object(), client=InferenceLLM())) as client:
        status = client.get('/api/status').json()
        assert status['max_temporary_upload_bytes'] == 500 * 1024 * 1024
        session = start(client)
        response = upload(client, session, 'dados.xlsx', raw)
        assert response.status_code == 200, response.text
        assert response.json()['datasets'][0]['preview'][0]['Valor'] == 25.5


def test_temporary_upload_accepts_exact_limit_and_rejects_one_byte_over(settings):
    settings.max_temporary_upload_bytes = 1024
    base = b'Nome,Quantidade\nTeste,2\n'
    raw = base + b'\n' * (1024 - len(base))
    llm = InferenceLLM()
    with TestClient(create_app(settings, mcp=object(), client=llm)) as client:
        session = start(client)
        assert upload(client, session, raw=raw).status_code == 200
        assert upload(client, session, raw=raw + b'\n').status_code == 413
        assert len(llm.calls) == 1


def expanded_workbook(megabytes=52):
    book = Workbook()
    book.active.append(['Valor'])
    book.active.append([25.5])
    buffer = io.BytesIO()
    book.save(buffer)
    with (ZipFile(buffer, 'a', compression=ZIP_DEFLATED) as archive,
          archive.open('large-unused.xml', 'w') as stream):
        stream.write(b'<root>')
        chunk = b' ' * (1024 * 1024)
        for _ in range(megabytes):
            stream.write(chunk)
        stream.write(b'</root>')
    return buffer.getvalue()


def test_temporary_xlsx_accepts_expansion_above_50_mb(settings):
    raw = expanded_workbook()
    assert len(raw) < 1024 * 1024
    with pytest.raises(ImportFailure, match='descompactado de 50 MB'):
        office_archive(raw)
    with TestClient(create_app(settings, mcp=object(), client=InferenceLLM())) as client:
        response = upload(client, start(client), 'estudo-sintetico.xlsx', raw)
        assert response.status_code == 200, response.text
        assert response.json()['datasets'][0]['preview'][0]['Valor'] == 25.5


def test_temporary_expansion_limit_reports_size_and_processing_stage(settings, caplog):
    settings.max_temporary_expanded_bytes = 1024 * 1024
    llm = InferenceLLM()
    with TestClient(create_app(settings, mcp=object(), client=llm)) as client:
        response = upload(client, start(client), 'estudo-sintetico.xlsx', expanded_workbook(2))
        assert response.status_code == 413
        assert 'descompactado de 1 MB' in response.json()['detail']
        assert response.headers['X-Import-Stage'] == 'read'
        assert 'stage=read status=413' in caplog.text
        assert 'estudo-sintetico' not in caplog.text
        assert llm.calls == []


def test_temporary_import_does_not_inherit_legacy_row_and_character_limits(settings):
    raw = ('Nome,Quantidade\n' + (('x' * 42) + ',2\n') * 50001).encode()
    assert len(raw) > 2_000_000
    assert settings.max_dataset_rows == 50000
    with TestClient(create_app(settings, mcp=object(), client=InferenceLLM())) as client:
        response = upload(client, start(client), raw=raw)
        assert response.status_code == 200, response.text
        assert response.json()['datasets'][0]['row_count'] == 50001


def test_streaming_xml_validation_rejects_declarations_across_chunk_boundaries():
    book = Workbook()
    book.active.append(['Valor'])
    book.active.append([1])
    buffer = io.BytesIO()
    book.save(buffer)
    with ZipFile(buffer, 'a', compression=ZIP_DEFLATED) as archive:
        archive.writestr('unsafe.xml', b' ' * (64 * 1024 - 4) + b'<!DOCTYPE x><x/>')
    with pytest.raises(ImportFailure, match='declarações não permitidas'):
        read_tables('dados.xlsx', buffer.getvalue(), 100)


@pytest.mark.parametrize('limit, expected', [('characters', 'caracteres'), ('rows', 'linhas')])
def test_temporary_content_limit_errors_are_specific(settings, limit, expected):
    if limit == 'characters':
        settings.max_temporary_characters = 1024
        raw = ('Nome,Quantidade\n' + ('x' * 100 + ',2\n') * 20).encode()
    else:
        settings.max_temporary_rows = 1
        raw = b'Nome,Quantidade\nA,1\nB,2'
    with TestClient(create_app(settings, mcp=object(), client=InferenceLLM())) as client:
        response = upload(client, start(client), raw=raw)
        assert response.status_code == 413
        assert expected in response.json()['detail']
