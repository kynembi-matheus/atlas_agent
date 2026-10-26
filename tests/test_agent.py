"""Testes de contrato de dados, execução, API e integração MCP simulada."""
import asyncio
import io
import json
from types import SimpleNamespace
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pydantic import SecretStr

from agents.orchestrator import Agent, RunManager
from core.config import Settings
from api.main import create_app
from core.models import AgentError, Run, RunRequest
from repositories.local_store import Store, now
from tools.data_tools import Toolbox, validate_sql


def new_run(status='running'):
    return Run(id=str(uuid4()), conversation_id=str(uuid4()), prompt='Analisar dados',
               status=status, created_at=now(), updated_at=now())


class FakeMCP:
    def __init__(self):
        self.calls = []

    async def discover(self):
        return ['list_databases', 'list_tables', 'describe_table', 'query']

    async def call(self, name, args):
        self.calls.append((name, args))
        if name == 'list_databases':
            return [{'database': 'cora'}, {'database': 'outro'}]
        if name == 'list_tables':
            return [{'table': 'vendas'}]
        if name == 'describe_table':
            return [{'column': 'valor', 'type': 'numeric'}]
        return {'rows': [{'equipe': 'A', 'valor': 10}, {'equipe': 'A', 'valor': 20},
                         {'equipe': 'B', 'valor': 5}]}


class ScriptLLM:
    def __init__(self, actions):
        self.responses = self
        self.actions = list(actions)
        self.inputs = []

    async def create(self, **kwargs):
        self.inputs.append(kwargs)
        action = self.actions.pop(0)
        if callable(action):
            action = action(kwargs)
        if isinstance(action, str):
            return SimpleNamespace(output=[], output_text=action)
        name, args = action
        return SimpleNamespace(output=[SimpleNamespace(type='function_call', name=name,
                               arguments=json.dumps(args), call_id=str(uuid4()))], output_text='')


@pytest.mark.parametrize('sql', [
    'DELETE FROM vendas', 'SELECT 1; SELECT 2',
    'WITH x AS (DELETE FROM vendas RETURNING *) SELECT * FROM x',
    'SELECT * INTO nova FROM vendas', 'SELECT * FROM vendas FOR UPDATE',
    "SELECT pg_read_file('/etc/passwd')", "SELECT nextval('sequencia')",
    "SELECT public.dblink_exec('x', 'DELETE FROM vendas')", 'COPY vendas TO STDOUT',
])
def test_rejects_unsafe_sql(sql):
    with pytest.raises(AgentError):
        validate_sql(sql)


@pytest.mark.parametrize('sql', [
    'SELECT equipe, SUM(valor) FROM vendas GROUP BY equipe',
    'WITH x AS (SELECT 1 AS valor) SELECT valor FROM x',
    'SELECT 1 AS valor UNION ALL SELECT 2',
])
def test_accepts_read_queries(sql):
    assert validate_sql(sql)


async def test_discovery_required_and_database_allowlist(settings, store):
    settings.allowed_databases = ['cora']
    mcp = FakeMCP()
    tools = Toolbox(settings, store, mcp, new_run())
    with pytest.raises(AgentError, match='list_databases'):
        await tools.call('query', {'database': 'cora', 'sql': 'SELECT 1'})
    assert await tools.call('list_databases', {}) == {'databases': ['cora']}
    with pytest.raises(AgentError):
        await tools.call('list_tables', {'database': 'outro'})
    await tools.call('list_tables', {'database': 'cora', 'schema': 'public'})
    assert mcp.calls[-1][1] == {'database': 'cora', 'schema': 'public'}


async def test_query_limits_and_lineage(settings, store):
    mcp = FakeMCP()
    tools = Toolbox(settings, store, mcp, new_run())
    await tools.call('list_databases', {})
    result = await tools.call('query', {'database': 'cora', 'sql': 'SELECT equipe, valor FROM vendas', 'max_rows': 100000})
    assert mcp.calls[-1][1]['max_rows'] == 10
    dataset = store.get_dataset(result['dataset_id'])
    assert dataset['source']['database'] == 'cora'
    assert dataset['source']['sql'].startswith('SELECT')
    assert len(dataset['rows']) == 3


async def test_transform_stats_and_original_preserved(settings, store):
    original = store.save_dataset([{'equipe': 'A', 'valor': 10}, {'equipe': 'A', 'valor': 20},
                                   {'equipe': 'B', 'valor': None}], {'kind': 'test'})
    tools = Toolbox(settings, store, FakeMCP(), new_run())
    result = await tools.call('transform_dataset', {'dataset_id': original['id'], 'operations': [
        {'op': 'fill_nulls', 'columns': ['valor'], 'value': 0},
        {'op': 'group', 'columns': ['equipe'], 'aggregations': {'valor': 'sum'}},
        {'op': 'sort', 'columns': ['valor'], 'ascending': False},
    ]})
    transformed = store.get_dataset(result['dataset_id'])
    assert transformed['rows'] == [{'equipe': 'A', 'valor': 30.0}, {'equipe': 'B', 'valor': 0.0}]
    assert transformed['source']['parent_dataset_id'] == original['id']
    assert store.get_dataset(original['id'])['rows'][2]['valor'] is None
    stats = await tools.call('analyze_dataset', {'dataset_id': result['dataset_id']})
    assert stats['statistics']['valor']['mean'] == 15


async def test_exports_neutralize_formulas_and_report_escapes_html(settings, store):
    dataset = store.save_dataset([{'nome': '=HYPERLINK("bad")', 'valor': 3},
                                  {'nome': '<script>alert(1)</script>', 'valor': -2}], {'kind': 'test'})
    run = new_run()
    tools = Toolbox(settings, store, FakeMCP(), run)
    xlsx = await tools.call('export_dataset', {'dataset_id': dataset['id'], 'format': 'xlsx'})
    book = load_workbook(io.BytesIO((store.artifacts_dir / xlsx['filename']).read_bytes()))
    assert book['Dados']['A2'].value.startswith("'=")
    assert book['Dados']['A2'].data_type == 's'
    assert book['Dados']['B3'].value == -2
    csv = await tools.call('export_dataset', {'dataset_id': dataset['id'], 'format': 'csv'})
    assert "'=HYPERLINK" in (store.artifacts_dir / csv['filename']).read_text(encoding='utf-8-sig')
    result = await tools.call('create_report', {'title': '<script>bad</script>', 'body': '<img src=x onerror=bad>',
                                              'dataset_ids': [dataset['id']], 'chart_x': 'nome', 'chart_y': 'valor'})
    html = (store.artifacts_dir / result['artifacts'][1]['filename']).read_text(encoding='utf-8')
    assert '<script>' not in html and '<img src=x' not in html
    assert '&lt;script&gt;' in html and 'magnitude' in html
    assert len(run.artifacts) == 4


async def test_docs_search_and_traversal_blocked(settings, store, tmp_path):
    settings.docs_dir.mkdir()
    (settings.docs_dir / 'regra.md').write_text('A métrica vendas usa o mês de competência.', encoding='utf-8')
    (tmp_path / 'secret.txt').write_text('segredo')
    tools = Toolbox(settings, store, FakeMCP(), new_run())
    result = await tools.call('search_docs', {'query': 'competência'})
    assert result[0]['path'] == 'docs/regra.md'
    with pytest.raises(AgentError):
        await tools.call('read_doc', {'path': '../secret.txt'})
    with pytest.raises(AgentError):
        await tools.call('read_doc', {'path': str(tmp_path / 'secret.txt')})


async def test_agent_full_tool_loop_generates_downloads(settings, store):
    dataset = store.save_dataset([{'mes': 'jan', 'valor': 10}, {'mes': 'fev', 'valor': 30}], {'kind': 'test'})
    llm = ScriptLLM([
        ('update_plan', {'steps': ['Calcular estatísticas', 'Gerar relatório e planilha']}),
        ('analyze_dataset', {'dataset_id': dataset['id']}),
        ('create_report', {'title': 'Resumo mensal', 'body': 'Total 40, média 20; dois meses de exemplo.',
                           'dataset_ids': [dataset['id']], 'chart_x': 'mes', 'chart_y': 'valor'}),
        ('export_dataset', {'dataset_id': dataset['id'], 'format': 'xlsx'}),
        'Relatório gerado com média 20. Os arquivos estão disponíveis para download.',
    ])
    run = new_run()
    store.save_run(run)
    manager = RunManager(Agent(settings, store, FakeMCP(), llm))
    request = RunRequest(prompt='Gere um relatório.', dataset_ids=[dataset['id']])
    submitted = manager.submit(request)
    await manager.tasks[submitted.id]
    finished = store.get_run(submitted.id)
    assert finished.status == 'completed'
    assert len(finished.artifacts) == 3
    assert len([e for e in finished.events if e['kind'] == 'tool_result']) == 4
    assert all(not call['store'] for call in llm.inputs)
    assert all('reasoning.encrypted_content' in call['include'] for call in llm.inputs)


async def test_tool_failure_can_be_corrected(settings, store):
    llm = ScriptLLM([('query', {'database': 'cora', 'sql': 'SELECT 1'}),
                    ('list_databases', {}), 'Os bancos disponíveis foram listados.'])
    run = new_run()
    agent = Agent(settings, store, FakeMCP(), llm)
    await agent.execute(run, RunRequest(prompt='Explore os bancos.'))
    assert any(e['kind'] == 'tool_error' for e in run.events)
    assert 'listados' in run.answer


async def test_cancel_running_and_queued(settings, store):
    started = asyncio.Event()
    class WaitingLLM:
        def __init__(self):
            self.responses = self
        async def create(self, **kwargs):
            started.set()
            await asyncio.Event().wait()
    settings.max_concurrent_runs = 1
    manager = RunManager(Agent(settings, store, FakeMCP(), WaitingLLM()))
    first = manager.submit(RunRequest(prompt='Esperar'))
    await started.wait()
    second = manager.submit(RunRequest(prompt='Outra tarefa'))
    assert (await manager.cancel(second.id)).status == 'cancelled'
    assert (await manager.cancel(first.id)).status == 'cancelled'
    await manager.close()


async def test_step_limit_and_conversation_lock(settings, store):
    settings.max_agent_steps = 2
    llm = ScriptLLM([('list_databases', {}), 'Descobri cora e outro; a análise ficou pendente.'])
    manager = RunManager(Agent(settings, store, FakeMCP(), llm))
    first = manager.submit(RunRequest(prompt='Explore'))
    with pytest.raises(AgentError, match='andamento'):
        manager.submit(RunRequest(prompt='Continuar', conversation_id=first.conversation_id))
    await manager.tasks[first.id]
    finished = store.get_run(first.id)
    assert finished.status == 'completed'
    assert 'limite de etapas atingido' in finished.answer
    assert 'análise ficou pendente' in finished.answer
    assert len(llm.inputs) == 2
    assert llm.inputs[-1]['tools'] == []


def test_agent_step_budget_is_at_most_60():
    from pydantic import ValidationError
    assert Settings(_env_file=None).max_agent_steps == 60
    with pytest.raises(ValidationError):
        Settings(_env_file=None, max_agent_steps=61)


async def test_final_step_does_not_execute_tools(settings, store):
    settings.max_agent_steps = 1
    mcp = FakeMCP()
    llm = ScriptLLM([('list_databases', {})])
    with pytest.raises(AgentError, match='síntese'):
        await Agent(settings, store, mcp, llm).execute(new_run(), RunRequest(prompt='Explore'))
    assert not mcp.calls


async def test_independent_tools_share_one_round(settings, store):
    def batch(_):
        return SimpleNamespace(output=[
            SimpleNamespace(type='function_call', name='update_plan',
                            arguments=json.dumps({'steps': ['Descobrir fontes']}), call_id='plan'),
            SimpleNamespace(type='function_call', name='list_databases', arguments='{}', call_id='db'),
        ], output_text='')

    class BatchLLM(ScriptLLM):
        async def create(self, **kwargs):
            if not self.inputs:
                self.inputs.append(kwargs)
                return batch(kwargs)
            return await super().create(**kwargs)

    llm = BatchLLM(['Bancos descobertos.'])
    run = new_run()
    await Agent(settings, store, FakeMCP(), llm).execute(run, RunRequest(prompt='Explore'))
    assert len(llm.inputs) == 2
    assert llm.inputs[0]['parallel_tool_calls'] is True
    assert len([e for e in run.events if e['kind'] == 'tool_result']) == 2


def test_persistence_recovers_interrupted_runs(store):
    running, done = new_run(), new_run('completed')
    store.save_run(running)
    store.save_run(done)
    restarted = Store(store.root)
    restarted.recover()
    assert restarted.get_run(running.id).status == 'failed'
    assert restarted.get_run(done.id).status == 'completed'


def test_api_auth_upload_download_and_validation(settings):
    settings.app_api_key = SecretStr('application-secret')
    with TestClient(create_app(settings, mcp=FakeMCP(), client=ScriptLLM([]))) as client:
        assert client.get('/').status_code == 200
        assert client.get('/health').json() == {'status': 'ok'}
        assert client.get('/api/runs').status_code == 401
        client.headers['Authorization'] = 'Bearer application-secret'
        response = client.post('/api/datasets/upload', files={'file': ('dados.csv', 'equipe;valor\nA;10\nB;20\n'.encode())})
        assert response.status_code == 200, response.text
        assert response.json()['row_count'] == 2
        assert client.get('/api/datasets/' + response.json()['dataset_id']).status_code == 200
        assert client.post('/api/datasets/upload', files={'file': ('dados.xlsx', b'bad')}).status_code == 422
        assert client.get('/api/runs/not-a-uuid').status_code == 422
        assert client.post('/api/runs', json={'prompt': '  '}).status_code == 409
        assert client.get('/api/connections/pg-mcp').json()['connected']
        artifact = client.app.state.store.save_artifact(str(uuid4()), 'Teste', 'csv', b'valor\n1')
        assert client.get(artifact['url']).content == b'valor\n1'
        assert client.get(artifact['url']).headers['cache-control'] == 'no-store'
        del client.headers['Authorization']
        assert client.get(artifact['url']).status_code == 401


def test_missing_llm_key_is_actionable(settings):
    settings.openai_api_key = SecretStr('')
    with TestClient(create_app(settings, mcp=FakeMCP(), client=ScriptLLM([]))) as client:
        response = client.post('/api/runs', json={'prompt': 'Analise'})
        assert response.status_code == 409
        assert 'OPENAI' in response.json()['detail']


def test_oversized_upload_rejected(settings):
    settings.max_upload_bytes = 1024
    with TestClient(create_app(settings, mcp=FakeMCP(), client=ScriptLLM([]))) as client:
        assert client.post('/api/datasets/upload', files={'file': ('large.csv', b'x' * 1025)}).status_code == 413

async def test_spreadsheet_headers_are_also_neutralized(settings, store):
    dataset = store.save_dataset([{'=formula': '=data', 'normal': -20}], {'kind': 'test'})
    tools = Toolbox(settings, store, FakeMCP(), new_run())
    result = await tools.call('export_dataset', {'dataset_id': dataset['id'], 'format': 'xlsx'})
    book = load_workbook(io.BytesIO((store.artifacts_dir / result['filename']).read_bytes()))
    assert book['Dados']['A1'].value == "'=formula"
    assert book['Dados']['A1'].data_type == 's'
    assert store.get_dataset(dataset['id'])['rows'][0]['=formula'] == '=data'


def test_single_column_csv_has_no_false_delimiter(settings):
    with TestClient(create_app(settings, mcp=FakeMCP(), client=ScriptLLM([]))) as client:
        response = client.post('/api/datasets/upload', files={'file': ('dados.csv', b'valor\n10\n20\n')})
        assert response.status_code == 200
        assert response.json()['columns'] == ['valor']
        assert response.json()['preview'] == [{'valor': 10}, {'valor': 20}]


async def test_numeric_text_requires_explicit_conversion(settings, store):
    original = store.save_dataset([{'codigo': '001', 'valor': '10.50'},
                                   {'codigo': '002', 'valor': '20.25'},
                                   {'codigo': '003', 'valor': None}], {'kind': 'test'})
    tools = Toolbox(settings, store, FakeMCP(), new_run())
    stats = await tools.call('analyze_dataset', {'dataset_id': original['id']})
    assert set(stats['numeric_text_columns']) == {'codigo', 'valor'}
    assert stats['warnings']
    with pytest.raises(AgentError, match='to_numeric'):
        await tools.call('transform_dataset', {'dataset_id': original['id'], 'operations': [
            {'op': 'group', 'aggregations': {'valor': 'sum'}}]})
    converted = await tools.call('transform_dataset', {'dataset_id': original['id'], 'operations': [
        {'op': 'to_numeric', 'columns': ['valor']}]})
    assert converted['preview'][0] == {'codigo': '001', 'valor': 10.5}
    stats = await tools.call('analyze_dataset', {'dataset_id': converted['dataset_id']})
    assert stats['statistics']['valor']['mean'] == 15.375
    assert stats['null_counts']['valor'] == 1
    assert store.get_dataset(original['id'])['rows'][0]['valor'] == '10.50'


@pytest.mark.parametrize('value', ['1.234,56', 'invalid-private-value', 'inf', ''])
async def test_numeric_conversion_refuses_invalid_values(settings, store, value):
    original = store.save_dataset([{'valor': value}], {'kind': 'test'})
    tools = Toolbox(settings, store, FakeMCP(), new_run())
    with pytest.raises(AgentError, match='Conversão numérica recusada') as error:
        await tools.call('transform_dataset', {'dataset_id': original['id'], 'operations': [
            {'op': 'to_numeric', 'columns': ['valor']}]})
    if value:
        assert value not in str(error.value)
    with store.connect() as db:
        assert db.execute('SELECT COUNT(*) FROM datasets').fetchone()[0] == 1


async def test_sample_warning_survives_multiple_transformations_and_reports(settings, store):
    original = store.save_dataset([{'valor': 10}, {'valor': 20}],
                                  {'kind': 'postgresql', 'possibly_truncated': True})
    tools = Toolbox(settings, store, FakeMCP(), new_run())
    result = await tools.call('transform_dataset', {'dataset_id': original['id'], 'operations': [
        {'op': 'group', 'aggregations': {'valor': 'sum'}}]})
    result = await tools.call('transform_dataset', {'dataset_id': result['dataset_id'], 'operations': [
        {'op': 'select', 'columns': ['valor']}]})
    assert result['possibly_truncated'] and result['warnings']
    stats = await tools.call('analyze_dataset', {'dataset_id': result['dataset_id']})
    assert stats['scope'] == 'stored_rows_only' and stats['possibly_truncated']
    report = await tools.call('create_report', {'title': 'Amostra', 'body': 'Análise',
                                               'dataset_ids': [result['dataset_id']]})
    for artifact in report['artifacts']:
        assert 'não representam totais da população' in (
            store.artifacts_dir / artifact['filename']).read_text(encoding='utf-8')


async def test_pagination_and_null_contains(settings, store):
    original = store.save_dataset([{'nome': None}, {'nome': 'None'}, {'nome': 'outro'}], {'kind': 'test'})
    tools = Toolbox(settings, store, FakeMCP(), new_run())
    page = await tools.call('inspect_dataset', {'dataset_id': original['id'], 'offset': 1, 'limit': 1})
    assert page['offset'] == 1 and page['next_offset'] == 2 and page['preview_only']
    last = await tools.call('inspect_dataset', {'dataset_id': original['id'], 'offset': 2, 'limit': 1})
    assert last['next_offset'] is None and last['preview_only']
    filtered = await tools.call('transform_dataset', {'dataset_id': original['id'], 'operations': [
        {'op': 'filter', 'column': 'nome', 'comparison': 'contains', 'value': 'None'}]})
    assert filtered['preview'] == [{'nome': 'None'}]


async def test_tool_validation_reports_field_without_echoing_input(settings, store):
    tools = Toolbox(settings, store, FakeMCP(), new_run())
    with pytest.raises(AgentError, match='max_rows: int_parsing') as error:
        await tools.call('query', {'database': 'cora', 'sql': 'SELECT 1', 'max_rows': 'private-value'})
    assert 'private-value' not in str(error.value)
    with pytest.raises(AgentError, match='Informe column'):
        await tools.call('transform_dataset', {'dataset_id': 'unused', 'operations': [{'op': 'filter'}]})


async def test_agent_recovers_from_text_aggregation(settings, store):
    original = store.save_dataset([{'valor': '10'}, {'valor': '20'}], {'kind': 'test'})

    def sum_converted(kwargs):
        output = json.loads(kwargs['input'][-1]['output'])
        return ('transform_dataset', {'dataset_id': output['dataset_id'], 'operations': [
            {'op': 'group', 'aggregations': {'valor': 'sum'}}]})

    llm = ScriptLLM([
        ('transform_dataset', {'dataset_id': original['id'], 'operations': [
            {'op': 'group', 'aggregations': {'valor': 'sum'}}]}),
        ('transform_dataset', {'dataset_id': original['id'], 'operations': [
            {'op': 'to_numeric', 'columns': ['valor']}]}),
        sum_converted,
        'A soma das linhas armazenadas é 30.',
    ])
    run = new_run()
    await Agent(settings, store, FakeMCP(), llm).execute(run, RunRequest(prompt='Some os valores.'))
    assert any(e['kind'] == 'tool_error' and 'to_numeric' in e['message'] for e in run.events)
    results = [e['result'] for e in run.events if e['kind'] == 'tool_result']
    assert results[-1]['preview'] == [{'valor': 30}]
