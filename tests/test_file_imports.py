"""Anexos locais: dados, origem, limites e conteúdo ilegível, sem IA ou banco reais."""
import io
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen.canvas import Canvas
from test_agent import FakeMCP, ScriptLLM

from api.main import create_app
from integrations.file_imports import ImportFailure, import_file
from tools.data_tools import coverage


def package(parts):
    buffer = io.BytesIO()
    with ZipFile(buffer, 'w', ZIP_DEFLATED) as archive:
        archive.writestr('[Content_Types].xml', '<Types/>')
        for name, content in parts.items():
            archive.writestr(name, content)
    return buffer.getvalue()


def docx():
    return package({'word/document.xml': '''<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>
    <w:p><w:r><w:t>Receita no período: 100</w:t></w:r></w:p>
    <w:tbl><w:tr><w:tc><w:p><w:r><w:t>Valor</w:t></w:r></w:p></w:tc></w:tr>
    <w:tr><w:tc><w:p><w:r><w:t>100</w:t></w:r></w:p></w:tc></w:tr></w:tbl>
    </w:body></w:document>'''})


def pptx():
    # Slide IDs deliberately differ from document order.
    return package({
        'ppt/presentation.xml': '''<p:presentation xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><p:sldIdLst><p:sldId id="256" r:id="r2"/><p:sldId id="257" r:id="r1"/></p:sldIdLst></p:presentation>''',
        'ppt/_rels/presentation.xml.rels': '<Relationships><Relationship Id="r1" Target="slides/slide1.xml"/><Relationship Id="r2" Target="slides/slide2.xml"/></Relationships>',
        **{f'ppt/slides/slide{i}.xml': f'''<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><p:sp><p:txBody><a:p><a:r><a:t>Indicador {i}</a:t></a:r></a:p></p:txBody></p:sp><a:tbl><a:tr><a:tc><a:txBody><a:p><a:r><a:t>{i * 10}</a:t></a:r></a:p></a:txBody></a:tc></a:tr></a:tbl></p:sld>''' for i in (1, 2)},
    })


def xlsx():
    book = Workbook()
    sheet = book.active
    sheet.title = 'Vendas'
    sheet.append(['codigo', 'valor', '__aba_origem'])
    sheet.append(['001', 10, 'original'])
    sheet.append(['002', '=B2*2', 'original'])
    other = book.create_sheet('Custos')
    other.append(['codigo', 'valor'])
    other.append(['003', 5])
    buffer = io.BytesIO()
    book.save(buffer)
    book.close()
    return buffer.getvalue()


def pdf(blank=False):
    buffer = io.BytesIO()
    canvas = Canvas(buffer)
    canvas.drawString(50, 800, 'Receita: 100')
    canvas.showPage()
    if blank:
        canvas.showPage()
    canvas.save()
    return buffer.getvalue()


@pytest.mark.parametrize(('name', 'content', 'count'), [
    ('dados.csv', b'codigo;valor\n001;10\n002;20', 2),
    ('dados.xlsx', xlsx(), 3), ('relatorio.docx', docx(), 3),
    ('indicadores.pptx', pptx(), 4), ('relatorio.pdf', pdf(), 1),
], ids=['csv', 'xlsx', 'docx', 'pptx', 'pdf'])
def test_upload_formats_persist_analysis_data_and_source(settings, name, content, count):
    with TestClient(create_app(settings, mcp=FakeMCP(), client=ScriptLLM([]))) as client:
        response = client.post('/api/datasets/upload', files={'file': (name, content)})
        assert response.status_code == 200, response.text
        data = response.json()
        assert data['row_count'] == count
        assert data['source']['filename'] == name
        assert len(data['source']['sha256']) == 64
        assert data['source']['format'] == name.rsplit('.', 1)[1]
        stored = client.app.state.store.get_dataset(data['dataset_id'])
        assert len(stored['rows']) == count
        assert client.get('/api/datasets/' + data['dataset_id']).json()['source'] == data['source']


def test_xlsx_provenance_all_sheets_and_cached_formulas():
    rows, source = import_file('dados.XLSX', xlsx(), 100)
    assert rows[0]['codigo'] == '001'
    assert rows[0]['valor'] == 10
    assert rows[0]['__aba_origem'] == 'original'
    assert rows[0][source['sheet_column']] == 'Vendas'
    assert rows[-1][source['sheet_column']] == 'Custos'
    assert rows[-1][source['row_column']] == 2
    assert rows[1]['valor'] is None
    assert source['extraction_warnings']


def test_documents_keep_tables_and_slide_order():
    rows, _ = import_file('dados.docx', docx(), 100)
    assert rows[0]['texto'] == 'Receita no período: 100'
    assert rows[-1]['coluna_1'] == '100'
    assert rows[-1]['bloco'] == 2
    rows, _ = import_file('dados.pptx', pptx(), 100)
    assert rows[0]['texto'] == 'Indicador 2' and rows[0]['slide'] == 1
    assert rows[1]['coluna_1'] == '20' and rows[1]['slide'] == 1
    assert rows[2]['texto'] == 'Indicador 1' and rows[2]['slide'] == 2


def test_formula_only_row_without_cache_is_preserved():
    book = Workbook()
    book.active.append(['valor'])
    book.active.append(['=1+1'])
    buffer = io.BytesIO()
    book.save(buffer)
    book.close()
    rows, source = import_file('dados.xlsx', buffer.getvalue(), 100)
    assert len(rows) == 1 and rows[0]['valor'] is None
    assert rows[0][source['row_column']] == 2


def test_pdf_marks_pages_without_text_as_incomplete():
    rows, source = import_file('dados.pdf', pdf(blank=True), 100)
    assert rows[0]['pagina'] == 1 and '100' in rows[0]['texto']
    assert source['possibly_truncated']
    assert source['pages_without_text'] == [2]
    assert coverage({'original_source': source})['warnings']


def test_pdf_encrypted_and_image_only_rejected():
    writer = PdfWriter()
    writer.append(PdfReader(io.BytesIO(pdf())))
    writer.encrypt('test-password')
    buffer = io.BytesIO()
    writer.write(buffer)
    with pytest.raises(ImportFailure, match='senha'):
        import_file('dados.pdf', buffer.getvalue(), 100)
    buffer = io.BytesIO()
    canvas = Canvas(buffer)
    canvas.showPage()
    canvas.save()
    with pytest.raises(ImportFailure, match='OCR'):
        import_file('dados.pdf', buffer.getvalue(), 100)


@pytest.mark.parametrize('extension', ['csv', 'xlsx', 'docx', 'pptx', 'pdf'])
def test_limits_reject_without_saving_partial_dataset(extension):
    contents = {'csv': b'valor\n1\n2', 'xlsx': xlsx(), 'docx': docx(), 'pptx': pptx(), 'pdf': pdf()}
    with pytest.raises(ImportFailure) as error:
        import_file('dados.' + extension, contents[extension], 0)
    assert error.value.status == 413


@pytest.mark.parametrize('extension', ['csv', 'xlsx', 'docx', 'pptx', 'pdf'])
def test_invalid_files_have_safe_errors(settings, extension):
    with TestClient(create_app(settings, mcp=FakeMCP(), client=ScriptLLM([]))) as client:
        response = client.post('/api/datasets/upload', files={'file': ('dados.' + extension, b'private-content')})
        assert response.status_code == 422
        assert 'private-content' not in response.text


def test_zip_traversal_and_entity_declarations_rejected():
    for content in [package({'../bad': 'x'}), package({'word/document.xml': '<!DOCTYPE x [<!ENTITY e "bad">]><x/>'})]:
        with pytest.raises(ImportFailure):
            import_file('dados.docx', content, 100)


def test_unsupported_and_empty_files(settings):
    with TestClient(create_app(settings, mcp=FakeMCP(), client=ScriptLLM([]))) as client:
        assert client.post('/api/datasets/upload', files={'file': ('dados.exe', b'bad')}).status_code == 415
        assert client.post('/api/datasets/upload', files={'file': ('dados.csv', b'valor\n')}).status_code == 422
