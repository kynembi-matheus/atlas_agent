"""Extração local e limitada de anexos para os datasets de análise."""
import csv
import hashlib
import io
import posixpath
import re
from itertools import zip_longest
from pathlib import PurePosixPath
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pandas as pd
from openpyxl import load_workbook
from pypdf import PdfReader

from tools.data_tools import rows_of

SUPPORTED = {"csv", "xlsx", "pdf", "docx", "pptx"}
MAX_TEXT = 2_000_000
MAX_PARTS = 1000
MAX_EXPANDED = 50 * 1024 * 1024
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


class ImportFailure(Exception):
    def __init__(self, message: str, status: int = 422):
        super().__init__(message)
        self.status = status


def xml(content: bytes):
    # Office XML does not need DTDs or entity declarations.
    markup = content.replace(b"\x00", b"").upper()
    if b"<!DOCTYPE" in markup or b"<!ENTITY" in markup:
        raise ImportFailure("XML com declarações não permitidas.")
    return ET.fromstring(content)


def office_archive(content: bytes, *, max_expanded_bytes: int = MAX_EXPANDED) -> ZipFile:
    archive = ZipFile(io.BytesIO(content))
    entries = archive.infolist()
    if len(entries) > 5000:
        archive.close()
        raise ImportFailure("Arquivo Office excede o limite de 5000 arquivos internos.", 413)
    if sum(entry.file_size for entry in entries) > max_expanded_bytes:
        archive.close()
        raise ImportFailure(f"Arquivo Office excede o limite de conteúdo descompactado de {max_expanded_bytes / 1024**2:g} MB.", 413)
    names = [entry.filename for entry in entries]
    if len(names) != len(set(names)) or any(
        entry.flag_bits & 1 or ".." in PurePosixPath(entry.filename).parts
        or entry.filename.startswith("/") or "\\" in entry.filename
        for entry in entries
    ):
        archive.close()
        raise ImportFailure("Arquivo Office inválido ou protegido.")
    if "[Content_Types].xml" not in names:
        archive.close()
        raise ImportFailure("O arquivo não é um documento Office válido.")
    return archive


def headers(values):
    result, used = [], set()
    for index, value in enumerate(values, 1):
        base = str(value).strip() if value is not None else f"coluna_{index}"
        base = base or f"coluna_{index}"
        name, suffix = base, 2
        while name in used:
            name = f"{base}_{suffix}"
            suffix += 1
        used.add(name)
        result.append(name)
    return result


class Collector:
    def __init__(self, max_rows):
        self.max_rows = max_rows
        self.rows = []
        self.characters = 0
        self.columns = set()

    def add(self, row):
        self.characters += sum(len(value) for value in row.values() if isinstance(value, str))
        self.columns.update(row)
        if len(self.rows) >= self.max_rows or len(self.columns) > 200 or self.characters > MAX_TEXT:
            raise ImportFailure("Conteúdo excede os limites de linhas, 200 colunas ou 2 milhões de caracteres.", 413)
        self.rows.append(row)

    def text(self, text, **location):
        text = text.strip()
        for offset in range(0, len(text), 4000):
            self.add({**location, "tipo": "texto", "trecho": offset // 4000 + 1,
                      "texto": text[offset:offset + 4000]})

    def table(self, table, namespace, **location):
        for index, tr in enumerate(table.findall(namespace + "tr"), 1):
            values = ["\n".join("".join(p.itertext()) for p in cell.findall(namespace + "p"))
                      if namespace == W else "\n".join(t.text or "" for t in cell.iter(A + "t"))
                      for cell in tr.findall(namespace + "tc")]
            if any(value.strip() for value in values):
                self.add({**location, "tipo": "tabela", "linha": index,
                          **{f"coluna_{i}": value for i, value in enumerate(values, 1)}})


def read_xlsx(content, collector, source):
    with office_archive(content) as archive:
        # Validate XML before openpyxl reads it (including declarations).
        for name in archive.namelist():
            if name.endswith((".xml", ".rels")):
                xml(archive.read(name))
    book = load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
    formulas = None
    try:
        formulas = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
        if len(book.worksheets) > MAX_PARTS:
            raise ImportFailure("Planilha excede o limite de abas.", 413)
        source["sheets"] = []
        source["extraction_warnings"] = [
            "Fórmulas não são recalculadas: valores usam o cache salvo no Excel; cache ausente resulta em nulo.",
            "Abas são reunidas por nome de coluna; filtre a aba de origem antes de agregar estruturas diferentes.",
        ]
        source["sheet_column"], source["row_column"] = "__aba_origem", "__linha_origem"
        # Pick provenance names that cannot overwrite an input column, across all sheets.
        sheet_headers = []
        for sheet in book.worksheets:
            sheet.reset_dimensions()
            first = next(sheet.iter_rows(values_only=True), ())
            if len(first) > 198:
                raise ImportFailure("XLSX excede o limite de 198 colunas de dados por aba.", 413)
            sheet_headers.append(headers(first))
        used = {name for names in sheet_headers for name in names}
        for key in ("sheet_column", "row_column"):
            while source[key] in used:
                source[key] += "_"
            used.add(source[key])
        for sheet, formula_sheet, names in zip(book.worksheets, formulas.worksheets, sheet_headers):
            formula_sheet.reset_dimensions()
            count = 0
            for number, (values, raw_values) in enumerate(zip_longest(
                sheet.iter_rows(min_row=2, values_only=True),
                formula_sheet.iter_rows(min_row=2, values_only=True), fillvalue=(),
            ), 2):
                if not any(value is not None for value in raw_values):
                    continue
                if len(raw_values) > len(names):
                    raise ImportFailure("Há dados além dos cabeçalhos da aba. Preencha os cabeçalhos no XLSX.")
                row = dict(zip(names, list(values) + [None] * (len(names) - len(values))))
                collector.add({**row, source["sheet_column"]: sheet.title, source["row_column"]: number})
                count += 1
            source["sheets"].append({"name": sheet.title, "row_count": count})
    finally:
        book.close()
        if formulas is not None:
            formulas.close()


def read_docx(content, collector, source):
    with office_archive(content) as archive:
        body = xml(archive.read("word/document.xml")).find(W + "body")
        if body is None:
            raise ImportFailure("DOCX sem corpo de documento.")
        for index, node in enumerate(body, 1):
            if node.tag == W + "p":
                collector.text("".join(t.text or "" for t in node.iter(W + "t")), bloco=index)
            elif node.tag == W + "tbl":
                collector.table(node, W, bloco=index)
    source["extraction_warnings"] = ["Somente texto e tabelas do corpo; imagens, gráficos, cabeçalhos e notas não são extraídos."]


def read_pptx(content, collector, source):
    with office_archive(content) as archive:
        presentation = xml(archive.read("ppt/presentation.xml"))
        relationships = {node.attrib["Id"]: node for node in xml(archive.read("ppt/_rels/presentation.xml.rels"))}
        slides = presentation.findall(".//" + P + "sldId")
        if len(slides) > MAX_PARTS:
            raise ImportFailure("Apresentação excede o limite de slides.", 413)
        source["slide_count"] = len(slides)
        for index, slide in enumerate(slides, 1):
            rel = relationships[slide.attrib[R + "id"]]
            if rel.attrib.get("TargetMode") == "External":
                raise ImportFailure("Referência externa de slide não permitida.")
            target = rel.attrib["Target"]
            path = posixpath.normpath(posixpath.join("ppt", target)) if not target.startswith("/") else target.lstrip("/")
            if not re.fullmatch(r"ppt/slides/slide\d+\.xml", path):
                raise ImportFailure("Referência de slide inválida.")
            root = xml(archive.read(path))
            for shape in root.iter(P + "sp"):
                collector.text("\n".join("".join(t.text or "" for t in p.iter(A + "t"))
                                         for p in shape.iter(A + "p")), slide=index)
            for number, table in enumerate(root.iter(A + "tbl"), 1):
                collector.table(table, A, slide=index, tabela=number)
    source["extraction_warnings"] = ["Somente texto e tabelas dos slides; imagens, gráficos, notas e objetos incorporados não são extraídos."]


def read_pdf(content, collector, source):
    if not content.startswith(b"%PDF-"):
        raise ImportFailure("O arquivo não é um PDF válido.")
    reader = PdfReader(io.BytesIO(content))
    if reader.is_encrypted:
        raise ImportFailure("PDF protegido por senha. Envie uma cópia sem proteção.")
    if len(reader.pages) > MAX_PARTS:
        raise ImportFailure("PDF excede o limite de páginas.", 413)
    empty_pages = []
    for index, page in enumerate(reader.pages, 1):
        stream = page.get_contents()
        if stream is not None and len(stream.get_data()) > MAX_EXPANDED:
            raise ImportFailure("Página PDF excede o limite de conteúdo descompactado.", 413)
        text = page.extract_text() or ""
        if not text.strip():
            empty_pages.append(index)
        collector.text(text, pagina=index)
    source["page_count"] = len(reader.pages)
    source["pages_without_text"] = empty_pages
    source["possibly_truncated"] = bool(empty_pages)
    source["extraction_warnings"] = ["Extração de texto sem OCR; tabelas do PDF não têm estrutura tabular garantida. Imagens e gráficos não são interpretados."]
    if empty_pages:
        source["extraction_warnings"].append("Há páginas sem texto extraível; o conteúdo importado pode estar incompleto.")


def import_file(filename: str, content: bytes, max_rows: int) -> tuple[list[dict], dict]:
    extension = filename.rsplit(".", 1)[-1].lower()
    if extension not in SUPPORTED:
        raise ImportFailure("Envie PDF, XLSX, CSV UTF-8, DOCX ou PPTX.", 415)
    source = {"kind": f"{extension}_upload", "filename": filename.replace("\\", "/").rsplit("/", 1)[-1][:200],
              "format": extension, "sha256": hashlib.sha256(content).hexdigest(),
              "extraction_warnings": []}
    collector = Collector(max_rows)
    try:
        if extension == "csv":
            text = content.decode("utf-8-sig")
            try:
                delimiter = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|").delimiter
            except csv.Error:
                delimiter = ","
            frame = pd.read_csv(io.BytesIO(content), sep=delimiter, engine="python", encoding="utf-8-sig", nrows=max_rows + 1)
            for row in rows_of(frame):
                collector.add(row)
        else:
            {"xlsx": read_xlsx, "docx": read_docx, "pptx": read_pptx, "pdf": read_pdf}[extension](content, collector, source)
    except ImportFailure:
        raise
    except Exception as exc:
        # Do not leak parser exceptions or file content to HTTP responses.
        raise ImportFailure(f"Arquivo {extension.upper()} inválido ou ilegível. Verifique o formato e a proteção por senha.") from exc
    if not collector.rows:
        raise ImportFailure("O arquivo precisa conter dados ou texto extraível. PDFs digitalizados precisam de OCR antes do envio.")
    source["row_count"] = len(collector.rows)
    return rows_of(pd.DataFrame(collector.rows)), source
