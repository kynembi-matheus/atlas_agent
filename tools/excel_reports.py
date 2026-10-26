"""Relatórios XLSX com abas de negócio, fontes resumidas e gráficos editáveis."""
import io
import re
from typing import Literal

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, ConfigDict, Field

from core.models import AgentError
from tools.report_sources import report_source, report_warning
from tools.report_tables import ReportColumn, ReportTable, table_data


class ExcelSheet(ReportTable):
    max_rows: int = Field(default=1000, ge=1, le=10000)
    columns: list[ReportColumn] = Field(min_length=1, max_length=20)
    chart: Literal['none', 'bar', 'line'] = 'none'
    category_key: str | None = None
    value_key: str | None = None


class ExcelReport(BaseModel):
    """Gere relatório Excel profissional com resumo, abas analíticas e gráficos editáveis.

    Leia excel_report/SKILL.md. Use datasets reais e colunas escolhidas com rótulos de negócio.
    Fontes internas são descritas como Banco de dados interno; metadados técnicos não são publicados.
    Não gera fórmulas arbitrárias. Cada gráfico mostra até 20 registros da aba e preserva o sinal.
    """
    model_config = ConfigDict(extra='forbid')
    title: str = Field(min_length=1, max_length=150)
    period: str = Field(min_length=1, max_length=150)
    summary: str = Field(min_length=1, max_length=6000)
    dataset_ids: list[str] = Field(min_length=1, max_length=8)
    sheets: list[ExcelSheet] = Field(min_length=1, max_length=6)


def safe_cell(value):
    if isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@')):
        return "'" + value
    return value


def build_excel(args: ExcelReport, datasets: list[dict], warnings: dict[str, list[str]]) -> bytes:
    by_id = {dataset['id']: dataset for dataset in datasets}
    if sum(min(len(by_id.get(sheet.dataset_id, {}).get('rows', [])), sheet.max_rows) for sheet in args.sheets) > 30000:
        raise AgentError('O relatório Excel aceita até 30000 linhas no conjunto das abas. Agregue ou exporte o dataset completo separadamente.')
    book = Workbook()
    summary = book.active
    summary.title = 'Resumo'
    summary.append([safe_cell(args.title)])
    summary.append(['Período', safe_cell(args.period)])
    summary.append(['Resumo executivo', safe_cell(args.summary)])
    for source in dict.fromkeys(report_source(dataset['source']) for dataset in datasets):
        summary.append(['Fonte', safe_cell(source)])
    for warning in dict.fromkeys(report_warning(w) for values in warnings.values() for w in values):
        summary.append(['Limitação', safe_cell(warning)])
    for spec in args.sheets:
        headers, _, caption = table_data(spec, by_id)
        dataset = by_id[spec.dataset_id]
        name = re.sub(r'[\\/*?:\[\]]', ' ', spec.title).strip().strip("'")[:31] or 'Indicadores'
        if name.casefold() in {name.casefold() for name in book.sheetnames}:
            name = f'{name[:25]} {len(book.sheetnames)}'
        sheet = book.create_sheet(name)
        sheet.append([safe_cell(spec.title)])
        sheet.append([safe_cell(spec.context)])
        sheet.append(['Fonte: ' + safe_cell(report_source(dataset['source']))])
        sheet.append([caption])
        sheet.append([safe_cell(header) for header in headers])
        for row in dataset['rows'][:spec.max_rows]:
            sheet.append([safe_cell(row[column.key]) if not isinstance(row[column.key], (dict, list)) else safe_cell(str(row[column.key]))
                          for column in spec.columns])
        sheet.freeze_panes = 'A6'
        sheet.auto_filter.ref = f'A5:{get_column_letter(len(headers))}{max(5, sheet.max_row)}'
        for index, column in enumerate(spec.columns, 1):
            sheet.column_dimensions[get_column_letter(index)].width = min(45, max(18, len(headers[index - 1]) + 4))
            for row in range(6, sheet.max_row + 1):
                cell = sheet.cell(row, index)
                if isinstance(cell.value, (int, float)):
                    cell.number_format = '#,##0.00"%"' if column.unit == '%' else '"R$" #,##0.00' if column.unit == 'R$' else '#,##0.00'
        if spec.chart != 'none':
            keys = [column.key for column in spec.columns]
            if spec.category_key not in keys or spec.value_key not in keys:
                raise AgentError('As colunas do gráfico Excel devem estar nas colunas selecionadas da aba.')
            selected = dataset['rows'][:min(spec.max_rows, 20)]
            if not selected or any(not isinstance(row[spec.value_key], (int, float)) or isinstance(row[spec.value_key], bool) for row in selected):
                raise AgentError('O gráfico Excel exige medidas numéricas e linhas disponíveis; converta números textuais explicitamente.')
            graphic = LineChart() if spec.chart == 'line' else BarChart()
            graphic.title = spec.title
            graphic.y_axis.title = headers[keys.index(spec.value_key)]
            graphic.x_axis.title = headers[keys.index(spec.category_key)]
            graphic.add_data(Reference(sheet, min_col=keys.index(spec.value_key) + 1, min_row=5, max_row=5 + len(selected)), titles_from_data=True)
            graphic.set_categories(Reference(sheet, min_col=keys.index(spec.category_key) + 1, min_row=6, max_row=5 + len(selected)))
            graphic.width, graphic.height = 22, 11
            sheet.add_chart(graphic, f'{get_column_letter(len(headers) + 2)}5')
            sheet.cell(sheet.max_row + 2, 1, f'Gráfico: primeiros {len(selected)} registros exibidos na aba; sem agregação implícita.')
        for cell in sheet[5]:
            cell.fill = PatternFill('solid', fgColor='183D35')
            cell.font = Font(color='FFFFFF', bold=True)
    summary.column_dimensions['A'].width = 24
    summary.column_dimensions['B'].width = 95
    summary.freeze_panes = 'A2'
    for sheet_index, sheet in enumerate(book):
        sheet.sheet_view.showGridLines = False
        sheet.row_dimensions[1].height = 32
        sheet.cell(1, 1).font = Font(size=16, bold=True, color='183D35')
        for row in sheet:
            for cell in row:
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        for row in range(2, min(4, sheet.max_row) + 1):
            sheet.row_dimensions[row].height = 42
        if sheet is not summary:
            for row in range(1, 5):
                sheet.merge_cells(start_row=row, start_column=1, end_row=row,
                                  end_column=max(2, len(args.sheets[sheet_index - 1].columns)))
    summary.row_dimensions[3].height = 90
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()
